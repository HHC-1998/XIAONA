"""Word 生成工具（05 详解 §6）：用 python-docx 把结构化内容渲染为美观模板 Word 文档。

模板约定（05 详解 §3/§4）：
- 零新增依赖（python-docx 已随多类型文件上传引入）；
- 封面区（title 非空时）→ 章节（heading/paragraphs/bullets/table）→ 页脚；
- 表格单元格内 URL 用 add_hyperlink 构造真实超链接（Word 内 Ctrl+点击跳转），
  同时单元格显示完整 URL 文本可复制；
- 图片经 httpx 下载嵌入，失败降级跳过。

产物落盘到 artifacts/{user_id}/{task_id}/{filename}，返回相对路径供下载接口使用。
"""
import asyncio
import json
import logging
import os
import re
import uuid
from datetime import datetime

logger = logging.getLogger("xiaona")

_INDIGO = "4F46E5"        # 表头深色底
_INDIGO_HEX = "FF4F46E5"
_ZEBRA_HEX = "FFF8FAFC"
_FONT_CN = "微软雅黑"
_LINK_COLOR = "0563C1"    # Word 标准超链接蓝色


def _safe_filename(name: str) -> str:
    base = re.sub(r"[\\/:*?\"<>|\s]", "_", name or "").strip("._") or f"文档_{uuid.uuid4().hex[:6]}"
    base = base[:60]
    if not base.lower().endswith(".docx"):
        base += ".docx"
    return base


def _safe_text(v) -> str:
    if v is None:
        return ""
    if not isinstance(v, str):
        try:
            v = json.dumps(v, ensure_ascii=False)
        except Exception:
            v = str(v)
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", v)


def _is_url(text: str) -> bool:
    return isinstance(text, str) and text.strip().startswith(("http://", "https://"))


# 图片扩展名正则（表格「图片」列识别用，2026-08-24 起）
_IMG_EXT_RE = re.compile(r"\.(?:jpe?g|png|gif|webp)(?:\?|$)", re.IGNORECASE)


def _is_image_url(text: str) -> bool:
    """识别单元格里的图片：`![alt](https://.../x.jpg)` 或纯图片 URL（.jpg/.png/.gif/.webp）。"""
    if not isinstance(text, str):
        return False
    t = text.strip()
    m = re.match(r"^!\[[^\]]*\]\((\S+)\)$", t)
    if m:
        return bool(_IMG_EXT_RE.search(m.group(1)))
    return bool(_is_url(t) and _IMG_EXT_RE.search(t))


def _extract_image_url(text: str) -> str:
    """从 `![alt](url)` 或纯 URL 里提取图片 URL。"""
    if not isinstance(text, str):
        return ""
    m = re.match(r"^!\[[^\]]*\]\((\S+)\)$", text.strip())
    return (m.group(1) if m else text).strip()


def _parse_markdown_table(md: str) -> tuple[list[str], list[list[str]]]:
    """把 markdown 表格原文解析为 (headers, rows)（2026-08-24 起，2026-09-01 去首尾空 cell）。

    跳过分隔行（|---），按 | 切列并 strip；图片单元格 `![alt](URL)` 原样保留。
    **去首尾空 cell**（2026-09-01 起）：`| | 评估维度 | 分析结论 | |` 这类行
    split 后首尾是空串（LLM 残留 markdown 管道符所致），剥离后列数才与表头一致。
    非表格输入返回空。
    """
    if not isinstance(md, str):
        return [], []
    lines = [ln.strip() for ln in md.splitlines() if ln.strip()]
    if len(lines) < 2:
        return [], []
    headers: list[str] = []
    rows: list[list[str]] = []
    for ln in lines:
        if not ln.startswith("|") or "|" not in ln[1:]:
            continue
        body = ln.strip("|")
        cells = [c.strip() for c in body.split("|")]
        # 去首尾空 cell（`| | a | b | |` → [a, b]）
        while cells and not cells[0]:
            cells.pop(0)
        while cells and not cells[-1]:
            cells.pop()
        if not cells:
            continue
        if all(re.match(r"^:?-{1,}:?$", c or "-") for c in cells):
            continue  # 分隔行（|---|）
        if not headers:
            headers = cells
        else:
            rows.append(cells)
    return headers, rows


# 文本表格/中文标题识别（2026-09-01 起：换行保真后，纯文本表格与中文编号标题按结构渲染）
_ZH_H1_RE = re.compile(
    r"^(?:[一二三四五六七八九十]+、|第[一二三四五六七八九十百零\d]+[章节部分]|（[一二三四五六七八九十]+）)"
)
_NUM_H2_RE = re.compile(r"^\d+\.\d+(?:\.\d+)?\s+\S")
# 表格表头关键词（表头行判定：含任一关键词 或 列数≥3；2026-09-01 扩充覆盖中文对照表：
# 组件/数据对象/对比项/访问地址/前置条件/交付物/维度/指标 等，让「评估维度 分析结论」式表头命中）
_TABLE_HEADER_KW = ("序号", "项目", "金额", "规格", "名称", "数量", "单价",
                    "说明", "内容", "配置", "维度", "指标", "结论", "费用", "小计",
                    "组件", "数据对象", "对比项", "访问地址", "前置条件", "交付物",
                    "访问方式", "安全要求", "典型场景", "评估维度", "包干金额",
                    "数据源", "关键词", "趋势词", "商品名", "排名", "评分", "评论")


def _split_cols(line: str, simple: bool = False) -> list[str]:
    """表格行分列：Tab 分隔优先（表头行常见），全角空格（　）次之，2+ 半角空格兜底。

    simple=True 时对单空格对齐行做「首列+中间+末列」3 列切分（报价表数据行：
    `1 NAS 主机 x86_64 成品 NAS，4 盘位 7,000` → [1, 说明, 7,000]），
    由调用方用数字校验防止正文段落误判。
    """
    if "\t" in line:
        cols = [c.strip() for c in line.split("\t") if c.strip()]
        if len(cols) >= 2:
            return cols
    if "　" in line:
        cols = [c.strip() for c in re.split(r"　+", line) if c.strip()]
        if len(cols) >= 2:
            return cols
    cols = [c.strip() for c in re.split(r"\s{2,}", line.strip()) if c.strip()]
    if len(cols) >= 2:
        return cols
    if simple:
        stripped = line.strip()
        first_sp = stripped.find(" ")
        last_sp = stripped.rfind(" ")
        if first_sp > 0:
            if first_sp < last_sp:
                # ≥2 空格：首列+中间+末列 3 列（报价表数据行：序号 说明 金额）
                return [stripped[:first_sp].strip(),
                        stripped[first_sp + 1:last_sp].strip(),
                        stripped[last_sp + 1:].strip()]
            # 恰 1 空格：首列+第二列 2 列（中文对照表数据行）
            head, tail = stripped[:first_sp].strip(), stripped[first_sp + 1:].strip()
            if head and tail and len(head) <= 15 and "。" not in head:
                return [head, tail]
    return []


def _is_num(x: str) -> bool:
    return bool(re.match(r"^\d[\d,.\-–—%]*$", x))


def _split_text_table(lines: list[str], start: int):
    """文本表格块探测（2026-09-01 起）：纯文本报价表格/中文对照表（无 | 管道符）按结构识别。

    规则（保守）：
    - 表头行：可分列（≥2 列，Tab/全角/2+ 空格/单空格首切）且（含表格关键词 或 列数≥3）
      且首/末列非数字；**拒绝含 `|` 的行**（markdown 表格交给 markdown 路径，防 `|`/空列/`---` 残留）
      与 `- `/`* ` 开头行（列表误判）；
    - 数据行：可分列且**列数与表头一致（容差 ±1）**即可——不再强制首列序号/末列金额，
      覆盖中文对照表（末列为中文说明，如「评估维度 分析结论」表下的「架构适配性 系统由…构成。」）；
      数字末列仍作对齐到表头末列的加分信号；
    - 数据行 ≥2 返回 (headers, rows, next_index)；不命中返回 None。
    表头前的行（如「（一）硬件费用　包干封顶价 HKD 13,000」）不吞进表格，
    留给主循环按标题/段落渲染。
    """
    j = start
    n = len(lines)
    header = None
    header_cols: list[str] = []
    header_is_money = False  # 表头含金额类词（金额/HKD/价格/费用/小计/单价）时，数据行末列须为数字
    rows: list[list[str]] = []
    while j < n:
        line = (lines[j] or "").strip()
        if not line:
            break
        if "|" in line or re.match(r"^\s*[-*]\s", line):
            break  # markdown 表格行 / 列表行：不是文本表格，交给主循环其它分支
        if header is None:
            cols_std = _split_cols(line)  # 标准分隔（Tab/全角/2+ 空格）
            cols = cols_std or _split_cols(line, simple=True)
            if not cols:
                break
            # 「列数≥3」仅认标准分隔切出的列（单空格 simple 切 3 列不算数——
            # 普通长句含 2 个空格就切 3 列，会误判表头）；含关键词仍算表头
            has_kw = any(kw in line for kw in _TABLE_HEADER_KW)
            is_hdr = has_kw or (bool(cols_std) and len(cols_std) >= 3)
            if is_hdr and not (_is_num(cols[0]) or _is_num(cols[-1])):
                header = cols
                header_cols = cols
                header_is_money = any(k in "".join(cols) for k in ("金额", "HKD", "价格", "费用", "小计", "单价"))
            else:
                break  # 第一个可分列的行不是表头：放弃探测（回落主循环，防吞标题/正文）
        else:
            # 标题行（一、/（一）/1.1）不吞进表格：立即收尾
            if _ZH_H1_RE.match(line) or _NUM_H2_RE.match(line):
                break
            if re.match(r"^\s*[-*]\s", line):
                break  # 列表行（防 `- ` 数据行误判成表格行）
            cols = _split_cols(line) or _split_cols(line, simple=True)
            if not cols:
                break
            if abs(len(cols) - len(header_cols)) > 1:
                break  # 列数与表头差太多：不像该表格的行
            if header_is_money and not _is_num(cols[-1]):
                break  # 报价表数据行末列须为金额（防表格后总结句混入）
            if len(cols) < len(header_cols):
                if _is_num(cols[-1]):
                    # 数字末列（报价表：序号 说明 金额）→ 对齐到表头末列（金额列）
                    rows.append([cols[0]] + cols[1:-1]
                                + [""] * (len(header_cols) - len(cols)) + [cols[-1]])
                else:
                    # 非数字（中文对照表）→ 顺排靠前，不足列补空
                    rows.append(cols + [""] * (len(header_cols) - len(cols)))
            elif len(cols) > len(header_cols):
                # 超过表头列数（单空格对齐行含多个空格时 3 列切过细）：前 ncols-1 列保留、
                # 多余列合并进最后一列（防丢内容）
                rows.append(cols[:len(header_cols) - 1]
                            + [" ".join(cols[len(header_cols) - 1:])])
            else:
                rows.append(cols)
        j += 1
    if header is not None and len(rows) >= 2:
        ncols = max(len(header_cols), max(len(r) for r in rows))
        header_out = header_cols + [""] * (ncols - len(header_cols))
        rows_out = [r + [""] * (ncols - len(r)) for r in rows]
        return header_out, rows_out, j
    return None


def _strip_md_symbols(text: str) -> str:
    """剥离文本里的 markdown 符号（**加粗**/*斜体*/`代码`/~~删除~~）。

    2026-09-01 起：直通渲染时原文可能残留 LLM 加的 markdown 标记，
    渲染到 Word/PDF 不应出现 `**`、`*` 等符号。
    """
    if not isinstance(text, str):
        return ""
    t = text.strip()
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"\*(.+?)\*", r"\1", t)
    t = re.sub(r"`([^`]*)`", r"\1", t)
    t = re.sub(r"~~(.+?)~~", r"\1", t)
    t = re.sub(r"!\[([^\]]*)\]\((\S+)\)", r"\1", t)  # 图片标记 → 保留 alt
    return t.strip()


def _smart_break_lines(content: str) -> str:
    """零换行/极少换行长文本按规则插换行（2026-09-01 方案 C，纯规则、保真、不调 LLM）。

    历史消息经旧前端单行 input 粘贴时换行被剥成一行 7800 字，规则块解析无从下手。
    此函数对**超长行（>100 字）**在「中文编号标题 / 数字编号标题 / 编号列表项 / 表格表头关键词行」
    前插换行，让 `_parse_markdown_blocks` 能识别标题/表格/段落。
    - 逐行处理：已有换行的文本不被破坏（短行原样保留，含 Tab 表格行）；
    - 压缩行内多空格但**保留单 Tab/单空格**（表格行分隔不丢）；
    - 已有合理换行（≥6 个）则原样返回。
    """
    if not content or content.count("\n") >= 6:
        return content
    out: list[str] = []
    for ln in content.splitlines():
        if not ln or len(ln) <= 100:
            out.append(ln)
            continue
        text = re.sub(r"[ \t　]{2,}", " ", ln).strip()
        if not text:
            out.append(ln)
            continue
        text = _BREAK_RE.sub(r"\n", text)  # 只插换行，不重复插入断点文本
        out.extend(x for x in text.split("\n"))
    return "\n".join(out)


# _smart_break_lines 的断点正则（中文标题 | 数字编号标题 | 编号列表项 | 表头关键词行）
_BREAK_RE = re.compile(
    r"(?<!\n)(?=("
    r"(?:[一二三四五六七八九十]+、|第[一二三四五六七八九十百零\d]+[章节部分]|（[一二三四五六七八九十]+）)"  # 中文标题
    r"|\d+(?:\.\d+){1,2}[\s、]"  # 数字编号标题（1.1 / 2.2）
    r"|\d+[.、)]\s+\S"  # 编号列表项（1. 系统部署…）
    r"|(?:序号|项目|金额|规格|名称|数量|单价|说明|内容|配置|维度|指标|结论|费用|小计|"
    r"组件|数据对象|对比项|访问地址|前置条件|交付物|访问方式|安全要求|典型场景|评估维度|包干金额)"
    r"\s*(?:[\t　]|$)"  # 表头关键词后跟 Tab/全角空格/行尾
    r"))"
)


def _parse_markdown_blocks(content: str) -> list[tuple]:
    """把 markdown 原文解析为结构化块列表（2026-09-01 起，content 原文渲染用）。

    返回 [(kind, ...)]：
    - ("heading", level, text)：`#`/`##`/`###` 标题行 + 中文/编号标题（一、/第X章/（一）→1，1.1→2）；
    - ("table", headers, rows)：markdown 表格块（复用 _parse_markdown_table，图片列保留）
      或纯文本表格块（_split_text_table，报价表格末列金额）；
    - ("bullet", text)：`- ` / `* ` 列表行；
    - ("para", text)：普通段落（含 `1. ` 编号行，按段落渲染）。

    2026-09-01 方案 C 改动：
    - 入口先过 `_smart_break_lines`（零换行旧消息语义分行）；
    - markdown 表格识别扩展为「连续 `|` 行」：首行 `|` 开头且次行也 `|` 开头即成表，
      不再强制 `---` 分隔行（LLM 输出/原文可能无分隔行）；
    - 文本剥离 markdown 符号（`**`/`*`/`` ` ``）。
    """
    blocks: list[tuple] = []
    if not content:
        return blocks
    content = _smart_break_lines(content)
    lines = content.splitlines()
    i, n = 0, len(lines)
    while i < n:
        stripped = (lines[i] or "").strip()
        if not stripped:
            i += 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            blocks.append(("heading", len(m.group(1)), _strip_md_symbols(m.group(2))))
            i += 1
            continue
        # 表格块：连续 | 行（首行 | 开头且次行也 | 开头；有 --- 分隔行或无均可）
        if stripped.startswith("|") and i + 1 < n and lines[i + 1].strip().startswith("|"):
            block = [lines[i].strip()]
            j = i + 1
            while j < n and lines[j].strip().startswith("|"):
                block.append(lines[j].strip())
                j += 1
            headers, rows = _parse_markdown_table("\n".join(block))
            if headers:
                blocks.append(("table", headers, rows))
            i = j
            continue
        # 文本表格块（2026-09-01 起：纯文本报价表格按结构渲染；含 | 的行已被上面拦截）
        t = _split_text_table(lines, i)
        if t:
            blocks.append(("table", t[0], t[1]))
            i = t[2]
            continue
        # 中文一级标题（2026-09-01 起：一、 / 第X章 / （一））
        if _ZH_H1_RE.match(stripped):
            blocks.append(("heading", 1, _strip_md_symbols(stripped)))
            i += 1
            continue
        # 编号二级标题（2026-09-01 起：1.1 短行且不以数字结尾，防表格行误判）
        if _NUM_H2_RE.match(stripped) and len(stripped) <= 40 \
                and not re.search(r"\d[\d,.]*$", stripped):
            blocks.append(("heading", 2, _strip_md_symbols(stripped)))
            i += 1
            continue
        if re.match(r"^\s*[-*]\s+", stripped):
            blocks.append(("bullet", _strip_md_symbols(re.sub(r"^\s*[-*]\s+", "", stripped))))
        elif re.match(r"^\s*\d+[.、)]\s+", stripped):
            blocks.append(("para", _strip_md_symbols(re.sub(r"^\s*\d+[.、)]\s+", "", stripped))))
        else:
            blocks.append(("para", _strip_md_symbols(stripped)))
        i += 1
    return blocks


def _default_doc_name(title: str, ext: str) -> str:
    """文件名语义化兜底（2026-08-24 起）：用 title 提炼，title 空则「文档_时间戳」。"""
    if title:
        return f"{title}_{datetime.now():%Y%m%d_%H%M%S}{ext}"
    return f"文档_{datetime.now():%Y%m%d_%H%M%S}{ext}"


def _set_run_font(run, size=10.5, bold=False, color=None):
    """设置 run 中文字体 + 字号 + 颜色（中文字体需同时设 eastAsia）。"""
    from docx.shared import Pt
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", _FONT_CN)


def _add_hyperlink(paragraph, url: str, text: str):
    """在段落内追加真实超链接（python-docx 无内置 API，用 oxml 构造 w:hyperlink）。

    同时显示完整 URL 文本（可复制），Word 内 Ctrl+点击可跳转（05 详解 §4）。
    """
    from docx.oxml import OxmlElement
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.shared import RGBColor

    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    # 蓝色 + 下划线（Word 标准超链接样式）
    c = OxmlElement("w:color")
    c.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", _LINK_COLOR)
    rPr.append(c)
    u = OxmlElement("w:u")
    u.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "single")
    rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _set_cell_shading(cell, hex_fill: str):
    """设置表格单元格底色（w:shd）。"""
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "clear")
    shd.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}color", "auto")
    shd.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill", hex_fill)
    tcPr.append(shd)


async def _download_images(images: list[dict], task_dir: str) -> list[dict]:
    if not images:
        return []
    import httpx
    out = []
    async with httpx.AsyncClient(timeout=15, follow_redirects=True) as client:
        for i, img in enumerate(images[:20]):
            url = (img or {}).get("url") or ""
            if not url.startswith(("http://", "https://")):
                continue
            try:
                resp = await client.get(url, headers={"User-Agent": "Mozilla/5.0 XIAONA"})
                if resp.status_code != 200:
                    continue
                ext = ".png" if resp.headers.get("content-type", "").startswith("image/png") else ".jpg"
                path = os.path.join(task_dir, f"_img_{i}{ext}")
                with open(path, "wb") as f:
                    f.write(resp.content)
                out.append({"path": path, "url": url, "caption": _safe_text((img or {}).get("caption"))})
            except Exception as e:
                logger.warning("docx image download failed: %s (%s)", url, e.__class__.__name__)
    return out


async def generate_docx(params: dict, ctx: dict | None = None) -> dict:
    """生成 Word(.docx)。params 与 generate_pdf 完全一致（04/05 同数据契约）。

    sections：[{heading?, paragraphs?, bullets?, table?}]；images：[{url, caption?}]；
    markdown_tables：[str] 可选的 markdown 表格原文（2026-08-24 起，executor 注入的
    会话榜单表格），由代码解析渲染、图片列直接嵌入，不依赖 LLM 转写；
    content：可选，markdown 原文（2026-09-01 起，用户粘贴的长文本直接渲染——
    标题/表格/bullets/段落自动解析，跳过 sections 校验，不依赖 LLM 手工转 JSON）。
    ctx: {"user_id", "task_id", "artifacts_dir"}。缺 ctx 时只校验数据不落盘。
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    # 防御性归一：LLM 输出截断/畸形时 args 可能不是 dict（2026-09-01 起）
    params = params if isinstance(params, dict) else {}
    sections = params.get("sections") or []
    if not isinstance(sections, list):
        sections = []
    content = _safe_text(params.get("content"))
    if not content and not sections:
        return {"ok": False, "error": "sections 不能为空：请提供文档章节内容"}
    title = _safe_text(params.get("title"))
    # LLM 可能在 title 里传字面 \n 换行（封面排版会崩），清洗为空格（2026-09-01 起）
    title = title.replace("\r", " ").replace("\n", " ")
    # title 自动提取（2026-09-01 起）：LLM 未传 title 时用 content 首个非空行，
    # 文件名不再回退「文档_时间戳」（剥引号与首尾空白，≤40 字）
    if not title and content:
        first_line = next((ln.strip() for ln in content.splitlines() if ln.strip()), "")
        title = first_line.strip("“”\"' 　")[:40]
    subtitle = _safe_text(params.get("subtitle"))
    date = _safe_text(params.get("date")) or datetime.now().strftime("%Y-%m-%d")
    source = _safe_text(params.get("source"))

    # 校验表格列数一致性
    for sec in sections:
        table = sec.get("table") if isinstance(sec, dict) else None
        if not table:
            continue
        headers = table.get("headers") or []
        rows = table.get("rows") or []
        if not headers or not isinstance(headers, list):
            return {"ok": False, "error": "表格 headers 不能为空"}
        for row in rows:
            if len(row) != len(headers):
                return {"ok": False, "error": f"表格列数与表头不一致（期望 {len(headers)} 列）"}

    if not ctx:
        return {"ok": True, "path": "", "size": 0, "title": title, "sections": len(sections)}

    user_id = ctx.get("user_id", "")
    task_id = ctx.get("task_id", "")
    artifacts_dir = ctx.get("artifacts_dir", "")
    task_dir = os.path.join(artifacts_dir, user_id, task_id)
    os.makedirs(task_dir, exist_ok=True)

    # 收集图片：独立 images 参数 + 表格「图片」列里 ![alt](url) + markdown_tables 的 URL（2026-08-24 起）
    img_urls: list[dict] = []
    seen: set[str] = set()
    for item in params.get("images") or []:
        u = (item or {}).get("url") or ""
        if u and u not in seen:
            seen.add(u)
            img_urls.append({"url": u, "caption": (item or {}).get("caption")})
    for sec in sections:
        # 非 dict 元素 / 无 table 键（纯文字 section）跳过，勿对 None 调 .get（2026-09-01 起）
        if not isinstance(sec, dict):
            continue
        table = sec.get("table") if isinstance(sec.get("table"), dict) else None
        if not table:
            continue
        for row in table.get("rows") or []:
            for c in row:
                if _is_image_url(c):
                    u = _extract_image_url(c)
                    if u and u not in seen:
                        seen.add(u)
                        img_urls.append({"url": u})
    for md in params.get("markdown_tables") or []:
        _h, md_rows = _parse_markdown_table(md)
        for row in md_rows:
            for c in row:
                if _is_image_url(c):
                    u = _extract_image_url(c)
                    if u and u not in seen:
                        seen.add(u)
                        img_urls.append({"url": u})
    images = await _download_images(img_urls, task_dir)
    img_by_url = {}
    for img in images:
        if img.get("url"):
            img_by_url[img["url"]] = img["path"]

    doc = Document()
    # 页面默认字体（正文）
    style = doc.styles["Normal"]
    style.font.name = _FONT_CN
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_CN)

    # 封面区
    if title:
        # 品牌色条
        bar = doc.add_table(rows=1, cols=1)
        bar.autofit = True
        cell = bar.cell(0, 0)
        cell.text = "XIAONA  ·  小娜智能文档"
        _set_cell_shading(cell, "4F46E5")
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                _set_run_font(r, size=9, color=RGBColor(0xFF, 0xFF, 0xFF))
        doc.add_paragraph()
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p_title.add_run(title)
        _set_run_font(r, size=24, bold=True, color=RGBColor(0x43, 0x38, 0xCA))
        if subtitle:
            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p_sub.add_run(subtitle)
            _set_run_font(r, size=10, color=RGBColor(0x6B, 0x72, 0x80))
        info = f"生成日期：{date}" + (f"　　数据来源：{source}" if source else "")
        p_info = doc.add_paragraph()
        p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p_info.add_run(info)
        _set_run_font(r, size=10, color=RGBColor(0x6B, 0x72, 0x80))
        doc.add_page_break()

    def render_table(headers: list[str], rows: list[list[str]], doc_ctx) -> None:
        """渲染表格（sections 表格与 markdown_tables 共用，2026-08-24 起）。
        「图片」列单元格为 ![alt](url) 时嵌入图片（add_picture 到单元格段落），
        下载失败/嵌入异常降级为原文本，不阻断。
        """
        if not headers:
            return
        t = doc_ctx.add_table(rows=1 + len(rows), cols=len(headers))
        t.style = "Table Grid"
        try:
            t.autofit = True  # Word 自动分配列宽（2026-09-01 起：长内容不再挤压变形）
        except Exception:
            pass
        img_cols = {j for j, h in enumerate(headers) if "图片" in h or "图" == h.strip()}
        for j, h in enumerate(headers):
            cell = t.cell(0, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER  # 表头居中（2026-09-01 起）
            r = p.add_run(h)
            _set_run_font(r, size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            _set_cell_shading(cell, "4F46E5")
        for i, row in enumerate(rows):
            for j, c in enumerate(row):
                cell = t.cell(i + 1, j)
                cell.text = ""
                p = cell.paragraphs[0]
                placed = False
                if j in img_cols and _is_image_url(c):
                    path = img_by_url.get(_extract_image_url(c))
                    if path:
                        try:
                            run = p.add_run()
                            run.add_picture(path, width=Cm(2.2))
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            placed = True
                        except Exception as e:
                            logger.warning("docx embed table image failed: %s", e.__class__.__name__)
                if not placed:
                    if _is_url(c):
                        _add_hyperlink(p, c.strip(), c.strip())
                    elif c:
                        r = p.add_run(c)
                        _set_run_font(r, size=9)
                if (i + 1) % 2 == 0:
                    _set_cell_shading(cell, "F8FAFC")
        doc_ctx.add_paragraph()

    # content 原文渲染（2026-09-01 起：用户粘贴的 markdown 原文直接渲染，
    # 表格块复用 _parse_markdown_table，图片列保留嵌入，不依赖 LLM 转写 sections）
    for block in _parse_markdown_blocks(content):
        kind = block[0]
        if kind == "heading":
            level = min(block[1], 3)
            h = doc.add_heading("", level=level)
            r = h.add_run(_safe_text(block[2]))
            _set_run_font(r, size={1: 15, 2: 13, 3: 11.5}[level], bold=True,
                          color=RGBColor(0x43, 0x38, 0xCA))
        elif kind == "table":
            render_table(block[1], block[2], doc)
        elif kind == "bullet":
            para = doc.add_paragraph(style="List Bullet")
            r = para.add_run(_safe_text(block[1]))
            _set_run_font(r, size=10.5)
        else:  # para
            para = doc.add_paragraph()
            r = para.add_run(_safe_text(block[1]))
            _set_run_font(r, size=10.5)

    # 章节（补充分析；sections 元素非 dict 时跳过，不抛异常）
    for raw in sections:
        sec = raw if isinstance(raw, dict) else {}
        heading = _safe_text(sec.get("heading"))
        if heading:
            h = doc.add_heading("", level=1)
            r = h.add_run(heading)
            _set_run_font(r, size=15, bold=True, color=RGBColor(0x43, 0x38, 0xCA))
        for p in sec.get("paragraphs") or []:
            para = doc.add_paragraph()
            r = para.add_run(_safe_text(p))
            _set_run_font(r, size=10.5)
        for b in sec.get("bullets") or []:
            para = doc.add_paragraph(style="List Bullet")
            r = para.add_run(_safe_text(b))
            _set_run_font(r, size=10.5)
        table = sec.get("table")
        if table:
            headers = [_safe_text(h) for h in table.get("headers") or []]
            rows = [[_safe_text(c) for c in row] for row in table.get("rows") or []]
            render_table(headers, rows, doc)

    # markdown 表格原文（2026-08-24 起：fetch_session_messages 提取、executor 注入，
    # 由代码直接渲染——图片列 100% 保留且列严格对齐，不依赖 LLM 转写）
    for md in params.get("markdown_tables") or []:
        h, md_rows = _parse_markdown_table(md)
        if h:
            render_table(h, md_rows, doc)

    # 图片章节
    for img in images:
        try:
            doc.add_picture(img["path"], width=Cm(8))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if img["caption"]:
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(img["caption"])
                _set_run_font(r, size=8.5, color=RGBColor(0x6B, 0x72, 0x80))
        except Exception as e:
            logger.warning("docx embed image failed: %s", e.__class__.__name__)

    # 页脚
    section = doc.sections[0]
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run("小娜出品")
    _set_run_font(r, size=8, color=RGBColor(0x6B, 0x72, 0x80))

    try:
        # 文件名语义化（2026-08-24 起）：filename 缺失或「文档_<时间戳>」占位名时用 title 提炼
        filename_in = params.get("filename") or ""
        if not filename_in or re.match(r"^文档_\d{8}_\d{6}$", filename_in.strip()):
            filename_in = _default_doc_name(title, "")
        filename = _safe_filename(filename_in or f"文档_{datetime.now():%Y%m%d_%H%M%S}.docx")
        path = os.path.join(task_dir, filename)
        doc.save(path)
        size = os.path.getsize(path)
        rel = os.path.relpath(path, artifacts_dir).replace("\\", "/")
        logger.info("generate_docx ok: %s (%d bytes)", rel, size)
        return {"ok": True, "path": rel, "size": size, "title": title, "sections": len(sections)}
    except Exception as e:
        logger.warning("generate_docx failed", exc_info=True)
        return {"ok": False, "error": f"Word 生成失败：{e.__class__.__name__}"}
    finally:
        for img in images:
            try:
                os.remove(img["path"])
            except OSError:
                pass
6